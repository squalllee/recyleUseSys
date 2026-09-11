"""Import one repairable part per numbered workbook sheet.

Requires openpyxl and pymssql. The command is a dry run unless --apply is used.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys
import unicodedata

import openpyxl
import pymssql


DEFAULT_MATERIAL_NO = "R0.AB.0007.WA"
DEFAULT_SYSTEM = "R"
DEFAULT_SUBSYSTEM = "AB"
NO_LOCATION_VALUES = {"", "無", "N/A", "NA", "NONE", "-"}


def read_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values

    for line in path.read_text(encoding="utf-8-sig").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_location(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value)).upper()
    text = re.sub(r"[\s車側]", "", text)
    if text in {"倉庫", "倉庫-WMS", "WMS"}:
        return "WMS"
    match = re.fullmatch(r"(\d{1,2})([AB])(\d)", text)
    if match:
        return f"{int(match.group(1)):02d}{match.group(2)}{match.group(3)}"
    return text


def is_no_location(value: object) -> bool:
    normalized = unicodedata.normalize("NFKC", clean_text(value)).upper()
    return normalized in NO_LOCATION_VALUES


def find_last_value(rows: list[list[object]], column_index: int) -> str:
    for row in reversed(rows):
        if column_index < len(row):
            value = clean_text(row[column_index])
            if value:
                return value
    return ""


def parse_workbook(path: Path, material_no: str) -> tuple[list[dict], list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    candidates: list[dict] = []
    skipped: list[dict] = []

    try:
        for number in range(1, 156):
            sheet_name = f"{number:03d}"
            device_id = f"{material_no}_{number:04d}"

            if sheet_name not in workbook.sheetnames:
                skipped.append({
                    "sheet": sheet_name,
                    "device_id": device_id,
                    "reason": "工作表不存在",
                })
                continue

            sheet = workbook[sheet_name]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            if not rows:
                skipped.append({
                    "sheet": sheet_name,
                    "device_id": device_id,
                    "reason": "工作表沒有資料",
                })
                continue

            headers = [clean_text(value) for value in rows[0]]
            header_map = {header: index for index, header in enumerate(headers) if header}
            required_headers = ["設備名稱", "設備序號", "裝上位置"]
            missing_headers = [header for header in required_headers if header not in header_map]
            if missing_headers:
                raise ValueError(
                    f"工作表 {sheet_name} 缺少欄位：{', '.join(missing_headers)}"
                )

            # Some sheets contain pre-numbered template rows whose only value
            # is the first "項次" cell. They are not maintenance records and
            # must not become the sheet's last data row.
            populated_rows = [
                row for row in rows[1:]
                if any(clean_text(value) for value in row[1:])
            ]
            if not populated_rows:
                skipped.append({
                    "sheet": sheet_name,
                    "device_id": device_id,
                    "reason": "工作表沒有維修記錄",
                })
                continue

            last_row = populated_rows[-1]
            location_index = header_map["裝上位置"]
            current_location = (
                clean_text(last_row[location_index])
                if location_index < len(last_row)
                else ""
            )
            if is_no_location(current_location):
                skipped.append({
                    "sheet": sheet_name,
                    "device_id": device_id,
                    "reason": "最後一筆有資料的紀錄沒有裝上位置",
                    "source_location": current_location or None,
                })
                continue

            device_name = find_last_value(populated_rows, header_map["設備名稱"])
            if not device_name:
                raise ValueError(f"工作表 {sheet_name} 找不到設備名稱")

            source_serial_number = find_last_value(
                populated_rows, header_map["設備序號"]
            )
            candidates.append({
                "sheet": sheet_name,
                "device_id": device_id,
                "device_name": device_name,
                "serial_number": sheet_name,
                "source_serial_number": source_serial_number or None,
                "source_location": current_location,
                "normalized_location": normalize_location(current_location),
            })
    finally:
        workbook.close()

    return candidates, skipped


def build_location_index(root_rows: list[dict]) -> dict[str, dict]:
    index: dict[str, dict] = {}
    for root in root_rows:
        for value in (root["DeviceID"], root["DeviceName"]):
            key = normalize_location(value)
            if not key:
                continue
            existing = index.get(key)
            if existing and existing["DeviceID"] != root["DeviceID"]:
                raise ValueError(
                    f"位置正規化後重複：{key} 同時對應 "
                    f"{existing['DeviceID']} 與 {root['DeviceID']}"
                )
            index[key] = root
    return index


def select_existing_rows(cursor, device_ids: list[str]) -> dict[str, dict]:
    if not device_ids:
        return {}
    placeholders = ", ".join(["%s"] * len(device_ids))
    cursor.execute(
        f"""
        SELECT DeviceID, DeviceName, MaterialNo, SerialNumber,
               CurrentLocationDeviceID, [System], [SubSystem]
        FROM dbo.RepairableDevices WITH (UPDLOCK, HOLDLOCK)
        WHERE DeviceID IN ({placeholders})
        """,
        tuple(device_ids),
    )
    return {row["DeviceID"]: row for row in cursor.fetchall()}


def expected_database_row(row: dict, material_no: str, system: str, subsystem: str) -> dict:
    return {
        "DeviceID": row["device_id"],
        "DeviceName": row["device_name"],
        "MaterialNo": material_no,
        "SerialNumber": row["serial_number"],
        "CurrentLocationDeviceID": row["current_location_device_id"],
        "System": system,
        "SubSystem": subsystem,
    }


def rows_equal(existing: dict, expected: dict) -> bool:
    return all(existing.get(key) == value for key, value in expected.items())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--env-file", type=Path, default=Path("api-server/.env"))
    parser.add_argument("--material-no", default=DEFAULT_MATERIAL_NO)
    parser.add_argument("--system", default=DEFAULT_SYSTEM)
    parser.add_argument("--subsystem", default=DEFAULT_SUBSYSTEM)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--skip-unmatched-locations", action="store_true")
    args = parser.parse_args()

    if not args.input.is_file():
        raise FileNotFoundError(args.input)

    candidates, skipped = parse_workbook(args.input, args.material_no)
    env = read_env(args.env_file)
    required_env = ["DB_SERVER", "DB_USER", "DB_PASSWORD", "DB_DATABASE"]
    missing_env = [name for name in required_env if not env.get(name)]
    if missing_env:
        raise ValueError(f"資料庫設定缺少：{', '.join(missing_env)}")

    connection = pymssql.connect(
        server=env["DB_SERVER"],
        user=env["DB_USER"],
        password=env["DB_PASSWORD"],
        database=env["DB_DATABASE"],
        autocommit=False,
    )

    inserted_rows = 0
    updated_rows = 0
    unchanged_rows = 0
    unmatched_rows: list[dict] = []
    collision_rows: list[dict] = []

    try:
        cursor = connection.cursor(as_dict=True)
        cursor.execute("SET XACT_ABORT ON")
        cursor.execute("""
            SELECT DeviceID, DeviceName
            FROM dbo.RepairableDevices
            WHERE CurrentLocationDeviceID IS NULL
        """)
        location_index = build_location_index(cursor.fetchall())

        matched_rows: list[dict] = []
        for row in candidates:
            location = location_index.get(row["normalized_location"])
            if location is None:
                unmatched_rows.append({
                    "sheet": row["sheet"],
                    "device_id": row["device_id"],
                    "source_location": row["source_location"],
                    "normalized_location": row["normalized_location"],
                    "reason": "找不到對應的 RepairableDevices 位置根節點",
                })
                continue
            matched_rows.append({
                **row,
                "current_location_device_id": location["DeviceID"],
            })

        if args.apply and unmatched_rows and not args.skip_unmatched_locations:
            raise ValueError(
                "有裝上位置無法對應根節點；請先修正位置，或使用 "
                "--skip-unmatched-locations 明確略過"
            )

        existing_rows = select_existing_rows(
            cursor, [row["device_id"] for row in matched_rows]
        )
        new_rows: list[dict] = []
        update_rows: list[dict] = []
        for row in matched_rows:
            expected = expected_database_row(
                row, args.material_no, args.system, args.subsystem
            )
            existing = existing_rows.get(row["device_id"])
            if existing is None:
                new_rows.append(expected)
            elif rows_equal(existing, expected):
                unchanged_rows += 1
            elif (
                existing.get("MaterialNo") == args.material_no
                and existing.get("SerialNumber") == row["serial_number"]
            ):
                update_rows.append(expected)
            else:
                collision_rows.append({
                    "device_id": row["device_id"],
                    "existing": existing,
                    "expected": expected,
                })

        if collision_rows:
            raise ValueError(
                f"{len(collision_rows)} 筆 DeviceID 已存在但內容不同，已取消匯入"
            )

        if args.apply and new_rows:
            cursor.executemany(
                """
                INSERT INTO dbo.RepairableDevices
                (
                    DeviceID, DeviceName, MaterialNo, SerialNumber,
                    CurrentLocationDeviceID, [System], [SubSystem]
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                [
                    (
                        row["DeviceID"], row["DeviceName"], row["MaterialNo"],
                        row["SerialNumber"], row["CurrentLocationDeviceID"],
                        row["System"], row["SubSystem"],
                    )
                    for row in new_rows
                ],
            )
            inserted_rows = len(new_rows)

        if args.apply and update_rows:
            cursor.executemany(
                """
                UPDATE dbo.RepairableDevices
                SET DeviceName = %s,
                    MaterialNo = %s,
                    SerialNumber = %s,
                    CurrentLocationDeviceID = %s,
                    [System] = %s,
                    [SubSystem] = %s,
                    UpdatedAt = SYSUTCDATETIME()
                WHERE DeviceID = %s
                """,
                [
                    (
                        row["DeviceName"], row["MaterialNo"], row["SerialNumber"],
                        row["CurrentLocationDeviceID"], row["System"],
                        row["SubSystem"], row["DeviceID"],
                    )
                    for row in update_rows
                ],
            )
            updated_rows = len(update_rows)

        if args.apply:
            connection.commit()
        else:
            connection.rollback()

        summary = {
            "mode": "apply" if args.apply else "dry-run",
            "workbook": str(args.input.resolve()),
            "material_no": args.material_no,
            "system": args.system,
            "subsystem": args.subsystem,
            "sheet_range": "001-155",
            "workbook_candidates": len(candidates),
            "matched_location_rows": len(matched_rows),
            "inserted_rows": inserted_rows,
            "would_insert_rows": len(new_rows) if not args.apply else 0,
            "updated_rows": updated_rows,
            "would_update_rows": len(update_rows) if not args.apply else 0,
            "unchanged_rows": unchanged_rows,
            "skipped_no_location_or_sheet": len(skipped),
            "skipped_details": skipped,
            "unmatched_location_rows": len(unmatched_rows),
            "unmatched_location_details": unmatched_rows,
        }
        print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
        return 0
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"匯入失敗：{error}", file=sys.stderr)
        raise
